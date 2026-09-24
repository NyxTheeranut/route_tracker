#!/usr/bin/env python3
"""
Pushes the 7-Eleven and RTR store master lists into the same Google Sheet used
for daily visit sync, as two read-only reference tabs ("7-Eleven Stores",
"RTR Stores") -- so the store list and the visit history live in one place.

Run this whenever Google_MyMap_7-Eleven Database NTB.xlsx or
Google_MyMap_RTR Database.xlsx change. Each run fully replaces those two tabs
(not an incremental append), so it's always safe to re-run.

index.html has no store data baked into it -- it fetches everything live from
the "7-Eleven Stores" / "RTR Stores" tabs this script writes to, scoped per
signed-in user by the Apps Script backend. So running this script is the whole
update: nothing else needs to be rebuilt for the page to pick up new stores.
"""
import json
import datetime
import urllib.request
import urllib.error
from pathlib import Path

import openpyxl

# Keep this in sync with DEFAULT_SYNC_URL in index.html -- if you change one,
# change the other.
SYNC_URL = "https://script.google.com/macros/s/AKfycbxJQOlyXKHuqOXLFMFxRUuJ3z-I50OL8kkPiMz0uOKY4DqjK0WdVOmJBd1aUTuWaWMilA/exec"

DASHBOARD_DIR = Path(__file__).resolve().parent.parent


def _find(name):
    """Locate a file that lives outside this repo, in the Dashboard folder.
    The source spreadsheets and the secret were moved into "Store Database/" and
    "Config/"; the bare Dashboard root is kept as a fallback so an older layout
    (or a copy left at the top level) still works. Returns the Store Database/
    path when nothing is found, so the error message names a sensible location."""
    for folder in (DASHBOARD_DIR / "Store Database", DASHBOARD_DIR / "Config", DASHBOARD_DIR):
        candidate = folder / name
        if candidate.exists():
            return candidate
    return DASHBOARD_DIR / "Store Database" / name


SEVEN_ELEVEN_XLSX = _find("Google_MyMap_7-Eleven Database NTB.xlsx")
RTR_XLSX = _find("Google_MyMap_RTR Database.xlsx")

# Must match the SYNC_SECRET Script Property set in the Apps Script project.
# Without this, syncStores had no auth at all -- anyone who found the
# (public, embedded-in-index.html) deployment URL could overwrite the entire
# store list with one request. Lives in a plain file in the Dashboard folder's
# Config/ directory, one level above this repo -- same as the xlsx sources, it's
# deliberately never committed: this script and its repo are public on GitHub
# Pages, and a secret checked into a public repo isn't a secret.
SYNC_SECRET_FILE = _find("sync_secret.txt")


def to_float(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def clean(v):
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):
        return v.isoformat()
    return str(v).strip()


def extract_seven_eleven(path):
    stores = []
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Store Database"]
    rows = list(ws.iter_rows(values_only=True))[1:]
    for r in rows:
        code, name, ae, cm, status, district, subdistrict, lat, lng = (
            r[1], r[2], r[3], r[4], r[5], r[6], r[7], r[8], r[9]
        )
        if code is None:
            continue
        lat_f, lng_f = to_float(lat), to_float(lng)
        if lat_f is None or lng_f is None:
            continue
        stores.append({
            "id": f"7-{int(code)}",
            "type": "7-11",
            "code": str(int(code)),
            "name": clean(name),
            "person": clean(ae),
            "cm": clean(cm),
            "status": clean(status),
            "district": clean(district),
            "subdistrict": clean(subdistrict),
            "address": "",
            "lat": round(lat_f, 6),
            "lng": round(lng_f, 6),
        })
    return stores


def extract_rtr(path):
    stores = []
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Sheet1"]
    rows = list(ws.iter_rows(values_only=True))[1:]
    for r in rows:
        (zone, cluster, code, rtype, name, status, owner, addr, prov, amphur,
         tumbon, lat, lng, rsr, cm, pbh) = r
        if code is None:
            continue
        lat_f, lng_f = to_float(lat), to_float(lng)
        if lat_f is None or lng_f is None:
            continue
        stores.append({
            "id": f"R-{int(code)}",
            "type": "RTR",
            "code": str(int(code)),
            "name": clean(name),
            "person": clean(rsr),
            "cm": clean(cm),
            "status": clean(status),
            "district": clean(amphur),
            "subdistrict": clean(tumbon),
            "address": clean(addr),
            "lat": round(lat_f, 6),
            "lng": round(lng_f, 6),
        })
    return stores


PROGRESS_TOTAL_STEPS = 3


def progress(step, label, width=28):
    filled = int(width * step / PROGRESS_TOTAL_STEPS)
    bar = "█" * filled + "░" * (width - filled)
    pct = int(100 * step / PROGRESS_TOTAL_STEPS)
    print(f"\n[{bar}] {pct:3d}%  Step {step}/{PROGRESS_TOTAL_STEPS}: {label}")


def main():
    if not SEVEN_ELEVEN_XLSX.exists():
        raise SystemExit(f"Not found: {SEVEN_ELEVEN_XLSX}")
    if not RTR_XLSX.exists():
        raise SystemExit(f"Not found: {RTR_XLSX}")
    if not SYNC_SECRET_FILE.exists():
        raise SystemExit(
            f"Not found: {SYNC_SECRET_FILE}\n"
            "Create it containing the same value as the SYNC_SECRET Script "
            "Property in the Apps Script project, with no extra whitespace."
        )
    sync_secret = SYNC_SECRET_FILE.read_text(encoding="utf-8").strip()

    progress(1, "Reading 7-Eleven store file")
    seven_eleven_stores = extract_seven_eleven(SEVEN_ELEVEN_XLSX)
    progress(2, "Reading RTR store file")
    stores = seven_eleven_stores + extract_rtr(RTR_XLSX)
    print(f"Parsed {len(stores)} stores "
          f"({sum(1 for s in stores if s['type']=='7-11')} 7-Eleven, "
          f"{sum(1 for s in stores if s['type']=='RTR')} RTR)")

    payload = json.dumps(
        {"action": "syncStores", "stores": stores, "secret": sync_secret}
    ).encode("utf-8")
    req = urllib.request.Request(
        SYNC_URL,
        data=payload,
        method="POST",
        headers={"Content-Type": "text/plain;charset=utf-8"},
    )
    progress(3, "Uploading to Google Sheet")
    print("Uploading to Google Sheet...")
    try:
        with urllib.request.urlopen(req, timeout=60) as res:
            body = res.read().decode("utf-8")
    except urllib.error.HTTPError as e:
        raise SystemExit(f"Upload failed: HTTP {e.code}\n{e.read().decode('utf-8', 'replace')[:500]}")
    except urllib.error.URLError as e:
        raise SystemExit(f"Upload failed: {e.reason}")

    try:
        result = json.loads(body)
    except json.JSONDecodeError:
        raise SystemExit(
            "Upload failed: response wasn't JSON (the Web App URL may need to be "
            "redeployed with \"Who has access: Anyone\").\n"
            f"First 300 chars of response:\n{body[:300]}"
        )

    if result.get("ok"):
        print(f"Done -- {result.get('count', len(stores))} stores synced to the sheet.")
    else:
        raise SystemExit(f"Upload failed: {result.get('error')}")


if __name__ == "__main__":
    main()
